# build_revised_report.py
# Comprehensive builder for PaperTrail_Project_Report_Revised.docx and PDF.

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

from report_helpers import set_cell_border, set_cell_margins, set_cell_shading
from chapters_data_shortened import get_chapters_shortened
from chapters_data_part3 import get_references
from tables_data import get_tables_dict

TEMPLATE_PATH = r"C:\Users\Nimal V\Downloads\Sample Project_Report.docx"
OUTPUT_DIR = r"c:\Users\Nimal V\OneDrive\Desktop\nimalprojects\paperTrial\frontend\documentation"
OUTPUT_DOCX = os.path.join(OUTPUT_DIR, "PaperTrail_Project_Report_Revised.docx")
OUTPUT_PDF = os.path.join(OUTPUT_DIR, "PaperTrail_Project_Report_Revised.pdf")

OFFICIAL_TITLE_L1 = "AN AI-POWERED RESEARCH PAPER COMPARISON AND"
OFFICIAL_TITLE_L2 = "EVIDENCE VERIFICATION SYSTEM USING"
OFFICIAL_TITLE_L3 = "RETRIEVAL-AUGMENTED GENERATION (RAG)"
OFFICIAL_TITLE_FULL = f"{OFFICIAL_TITLE_L1} {OFFICIAL_TITLE_L2} {OFFICIAL_TITLE_L3}"

RIGHT_MARGIN_TWIPS = 9000

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=Pt(3), space_after=Pt(3),
          line_spacing=1.3, bold=False, italic=False, font_size=Pt(12), font_name='Times New Roman'):
    p = doc.add_paragraph()
    p.alignment = align
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(int(space_before.pt * 20)))
    sp.set(qn('w:after'), str(int(space_after.pt * 20)))
    sp.set(qn('w:line'), str(int(line_spacing * 240)))
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)

    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = font_size
    return p

def add_chapter_separator_page(doc, chapter_num, title):
    """
    Creates a dedicated chapter opening page matching the reference image:
    CHAPTER - N
    TITLE
    centered vertically and horizontally.
    """
    p_sp = doc.add_paragraph()
    pPr_sp = p_sp._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '4400') # ~220 pt to push down to middle
    pPr_sp.append(sp)

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr1 = p1._element.get_or_add_pPr()
    sp1 = OxmlElement('w:spacing')
    sp1.set(qn('w:before'), '0')
    sp1.set(qn('w:after'), '120')
    pPr1.append(sp1)
    r1 = p1.add_run(f"CHAPTER - {chapter_num}")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(28)
    r1.bold = False

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p2._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '120')
    sp2.set(qn('w:after'), '0')
    pPr2.append(sp2)
    r2 = p2.add_run(title)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(28)
    r2.bold = False

    # Page break after separator page
    p_br = doc.add_paragraph()
    r_br = p_br.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r_br._r.append(br)

def add_chapter_heading(doc, chapter_num, title):
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_num._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '240')
    sp.set(qn('w:after'), '80')
    pPr.append(sp)
    r1 = p_num.add_run(f"CHAPTER {chapter_num}")
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(18)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p_title._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '60')
    sp2.set(qn('w:after'), '240')
    pPr2.append(sp2)
    r2 = p_title.add_run(title)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(16)

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '180')
    sp.set(qn('w:after'), '60')
    pPr.append(sp)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    return p

def add_code_block(doc, code_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '100')
    sp.set(qn('w:after'), '100')
    sp.set(qn('w:line'), '240')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '360')
    pPr.append(ind)
    r = p.add_run(code_str)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.0)
    return p

def add_figure(doc, image_path, fig_no, caption, width=Inches(4.6)):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_img._element.get_or_add_pPr()
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '140')
        sp.set(qn('w:after'), '60')
        pPr.append(sp)
        run = p_img.add_run()
        run.add_picture(image_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p_cap._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '40')
    sp2.set(qn('w:after'), '140')
    sp2.set(qn('w:line'), '280')
    sp2.set(qn('w:lineRule'), 'auto')
    pPr2.append(sp2)
    r_cap = p_cap.add_run(f"[{fig_no}: {caption}]")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10.5)
    r_cap.bold = False

def add_table(doc, tbl_data):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_cap._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '120')
    sp.set(qn('w:after'), '60')
    pPr.append(sp)
    r = p_cap.add_run(f"{tbl_data['num']}: {tbl_data['title']}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)

    headers = tbl_data['headers']
    rows = tbl_data['rows']
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    border_kwargs = dict(
        top=dict(sz=4, val='single', color='888888', space='0'),
        bottom=dict(sz=4, val='single', color='888888', space='0'),
        left=dict(sz=4, val='single', color='888888', space='0'),
        right=dict(sz=4, val='single', color='888888', space='0'),
        insideH=dict(sz=4, val='single', color='CCCCCC', space='0'),
        insideV=dict(sz=4, val='single', color='CCCCCC', space='0')
    )

    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.text = h_text
        set_cell_shading(cell, "EAEAEA")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        set_cell_border(cell, **border_kwargs)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.0)

    for r_idx, row in enumerate(rows):
        for col_idx, val in enumerate(row):
            cell = tbl.cell(r_idx+1, col_idx)
            cell.text = str(val)
            if r_idx % 2 == 1:
                set_cell_shading(cell, "F9F9F9")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            set_cell_border(cell, **border_kwargs)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 and len(str(val)) > 20 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.0)

    p_sp = doc.add_paragraph()
    p_spPr = p_sp._element.get_or_add_pPr()
    sp_post = OxmlElement('w:spacing')
    sp_post.set(qn('w:before'), '0')
    sp_post.set(qn('w:after'), '100')
    p_spPr.append(sp_post)

def make_entry_paragraph(prefix, title, page_str, is_chapter=False, is_sub=False, pos=RIGHT_MARGIN_TWIPS):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')

    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '40' if is_chapter else '15')
    sp.set(qn('w:after'), '40' if is_chapter else '15')
    sp.set(qn('w:line'), '240')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)

    if is_sub:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '360')
        pPr.append(ind)

    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), str(pos))
    tabs.append(tab)
    pPr.append(tabs)
    p.append(pPr)

    # Label & Title Run
    r1 = OxmlElement('w:r')
    r1Pr = OxmlElement('w:rPr')
    rf1 = OxmlElement('w:rFonts')
    rf1.set(qn('w:ascii'), 'Times New Roman')
    rf1.set(qn('w:hAnsi'), 'Times New Roman')
    r1Pr.append(rf1)
    sz1 = OxmlElement('w:sz')
    sz1.set(qn('w:val'), '21' if is_chapter else '19')
    r1Pr.append(sz1)
    if is_chapter:
        r1Pr.append(OxmlElement('w:b'))
    r1.append(r1Pr)

    t1 = OxmlElement('w:t')
    label = f"{prefix}   {title}" if prefix else title
    t1.text = f"{label}\t"
    t1.set(qn('xml:space'), 'preserve')
    r1.append(t1)
    p.append(r1)

    # Page Number Run
    r2 = OxmlElement('w:r')
    r2Pr = OxmlElement('w:rPr')
    rf2 = OxmlElement('w:rFonts')
    rf2.set(qn('w:ascii'), 'Times New Roman')
    rf2.set(qn('w:hAnsi'), 'Times New Roman')
    r2Pr.append(rf2)
    sz2 = OxmlElement('w:sz')
    sz2.set(qn('w:val'), '21' if is_chapter else '19')
    r2Pr.append(sz2)
    if is_chapter:
        r2Pr.append(OxmlElement('w:b'))
    r2.append(r2Pr)

    t2 = OxmlElement('w:t')
    t2.text = str(page_str)
    r2.append(t2)
    p.append(r2)

    return p

def update_front_matter(doc, toc_entries=None, lot_entries=None, lof_entries=None):
    print("Updating front matter...")

    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.25)
        sec.right_margin = Inches(1.0)
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.5)

    # Title Page
    p1 = doc.paragraphs[1]
    p1.text = ""
    r1 = p1.add_run(OFFICIAL_TITLE_L1)
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(17)

    p2 = doc.paragraphs[2]
    p2.text = ""
    r2 = p2.add_run(OFFICIAL_TITLE_L2)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(17)

    p3 = doc.paragraphs[3]
    p3.text = ""
    r3 = p3.add_run(OFFICIAL_TITLE_L3)
    r3.bold = True
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(17)

    doc.paragraphs[4].text = ""
    doc.paragraphs[5].text = ""

    # Bonafide Certificate
    p27 = doc.paragraphs[27]
    p27.text = ""
    r27_1 = p27.add_run("Certified that this project report \"")
    r27_1.font.name = 'Times New Roman'
    r27_1.font.size = Pt(13)
    r27_2 = p27.add_run(OFFICIAL_TITLE_FULL)
    r27_2.bold = True
    r27_2.font.name = 'Times New Roman'
    r27_2.font.size = Pt(13)
    r27_3 = p27.add_run("\" is the bonafide work of SUBASH KUMAR K (723722243087), SIVAPRATHEESH S (723722243081), MOHAN KUMAR S (723722243302), and SURESH (723722243088) who carried out the project work under my supervision.")
    r27_3.font.name = 'Times New Roman'
    r27_3.font.size = Pt(13)

    # Acknowledgement: Tighten spacing so it fits on 1 page cleanly!
    doc.paragraphs[48].text = "ACKNOWLEDGEMENT"
    doc.paragraphs[49].text = "First, we extend our heartfelt gratitude to the management of V.S.B. College of Engineering Technical Campus for providing us with the state-of-the-art computational infrastructure, high-speed research facilities, and an intellectually inspiring learning environment to pursue this project."
    doc.paragraphs[51].text = "We express our sincere thanks to our beloved Principal, Dr. V. Velmurugan, M.E., Ph.D., for his continuous encouragement, leadership, and visionary academic guidance throughout our degree program."
    doc.paragraphs[53].text = "We also express our gratitude to the Vice Principal, Dr. R. Poornachandran, M.E., Ph.D., for his academic support, administrative cooperation, and regular motivation."
    doc.paragraphs[55].text = "We are highly grateful to Dr. R. Murugadoss, M.E., Ph.D., and Mrs. G. Nithya, M.E., (Ph.D), Head of the Department of Artificial Intelligence and Data Science, for their consistent encouragement, domain guidance, and constructive suggestions during technical seminars."
    doc.paragraphs[59].text = "We would like to express our deepest gratitude to our esteemed project guide, Mr. S. Soundhar, M.E., Assistant Professor, Department of Artificial Intelligence and Data Science, for his invaluable technical guidance, meticulous review of system architecture, and insightful recommendations throughout the lifecycle of this project."
    doc.paragraphs[61].text = "Finally, we extend our sincere gratitude to our parents, faculty members, lab technicians, and friends who supported us morally and technically in bringing this research project to successful completion."

    for pi in range(48, 62):
        p = doc.paragraphs[pi]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '40')
        sp.set(qn('w:after'), '40')
        sp.set(qn('w:line'), '240')
        sp.set(qn('w:lineRule'), 'auto')

    # POs & Profile: tighten line spacing so they don't have orphan spills
    for pi in range(74, 159):
        p = doc.paragraphs[pi]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '20')
        sp.set(qn('w:after'), '20')
        sp.set(qn('w:line'), '230')
        sp.set(qn('w:lineRule'), 'auto')

    # Abstract: P 159, 172, 174 (tighten so it fits cleanly on 1 page!)
    p159 = doc.paragraphs[159]
    p159.text = OFFICIAL_TITLE_FULL
    for r in p159.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    p172 = doc.paragraphs[172]
    p172.text = (
        "The exponential expansion of peer-reviewed scientific literature poses a formidable cognitive barrier for researchers, "
        "graduate scholars, and peer reviewers. Traditional academic search engines index keywords but provide zero cross-document "
        "relational synthesis. Conversely, general-purpose Large Language Models (LLMs) frequently generate hallucinated citations, "
        "phantom experimental results, and superficial summaries that lack verifiable grounding. This project presents an enterprise-grade "
        "AI-Powered Research Paper Comparison and Evidence Verification System using Retrieval-Augmented Generation (RAG) designed to perform "
        "automated, high-precision multi-document comparative analysis with absolute provenance. The framework integrates an automated PDF extraction pipeline "
        "using PyMuPDF that preserves exact 1-indexed page coordinates, a metadata-preserving sliding window chunker (512 tokens with "
        "100-token overlap), and a hybrid retrieval engine combining dense vector embeddings (SentenceTransformers all-MiniLM-L6-v2) "
        "with sparse BM25 lexical indexing fused via Reciprocal Rank Fusion (k=60). Synthesized claims are validated through Natural Language "
        "Inference (NLI) entailment scoring, generating clickable, page-level citations formatted as [Paper Title, p. X]. The system automatically "
        "constructs multi-dimensional comparative matrices spanning objectives, methodologies, architectures, datasets, metrics, and limitations. "
        "Furthermore, an automated metric consistency validation module audits reported benchmarks, identifying discrepancies in evaluation splits "
        "and measurement units. An enterprise security module utilizing Spring Security 6, stateless JWT, and BCrypt hashing enforces multi-tenant "
        "isolation and resource ownership verification. Experimental evaluation across 50 peer-reviewed papers demonstrated a Top-5 retrieval precision "
        "of 96.8%, citation accuracy of 97.4%, and an end-to-end processing latency of under 8 seconds for 3-paper comparison with 100% security test compliance."
    )
    for r in p172.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)

    p174 = doc.paragraphs[174]
    p174.text = (
        "Retrieval-Augmented Generation (RAG), scientific document intelligence, multi-paper comparison matrix, "
        "evidence grounding, page-level citations, metric consistency auditing, hallucination mitigation, "
        "Spring Boot, Spring Security 6, JWT authentication, PyMuPDF, FastAPI, hybrid vector search."
    )
    for r in p174.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)

    for pi in [159, 172, 174]:
        p = doc.paragraphs[pi]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '30')
        sp.set(qn('w:after'), '30')
        sp.set(qn('w:line'), '230')
        sp.set(qn('w:lineRule'), 'auto')

    # Tables replacement for TOC, LOT, LOF
    body = doc._body._element

    if toc_entries and lot_entries and lof_entries:
        # Table 2: TOC
        tbl2 = doc.tables[2]._element
        for item in toc_entries:
            p_el = make_entry_paragraph(item[0], item[1], item[2], item[3], item[4])
            tbl2.addprevious(p_el)
        body.remove(tbl2)

        # Table 3: LOT
        tbl3 = doc.tables[2]._element
        for item in lot_entries:
            p_el = make_entry_paragraph(item[0], item[1], item[2], False, False)
            tbl3.addprevious(p_el)
        body.remove(tbl3)

        # Table 4: LOF
        tbl4 = doc.tables[2]._element
        for item in lof_entries:
            p_el = make_entry_paragraph(item[0], item[1], item[2], False, False)
            tbl4.addprevious(p_el)
        body.remove(tbl4)

print("Front matter updater defined.")
