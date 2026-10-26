"""
Deep inspection of the Word document - look at all paragraphs around chapters
and understand the full structure.
"""
from docx import Document
from docx.oxml.ns import qn

doc_path = r"c:\Users\Nimal V\OneDrive\Desktop\nimalprojects\paperTrial\frontend\documentation\PaperTrail_Project_Documentation.docx"
doc = Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")

print("\n=== ALL PARAGRAPHS (from 240 onward - main content) ===")
for i, para in enumerate(doc.paragraphs[240:]):
    idx = i + 240
    text = para.text.strip()
    style = para.style.name
    # Check for page breaks in runs
    has_page_break = False
    for run in para.runs:
        if run._element.xml.find('w:lastRenderedPageBreak') >= 0 or run._element.xml.find('w:br') >= 0:
            has_page_break = True
            break
    # Check paragraph XML for page break
    if para._element.xml.find('w:pageBreak') >= 0 or 'w:type="page"' in para._element.xml:
        has_page_break = True
    
    pb_flag = " [PAGE_BREAK]" if has_page_break else ""
    if text:
        print(f"[{idx:4d}] [{style:25s}]{pb_flag} {text[:100]}")
    elif has_page_break:
        print(f"[{idx:4d}] [{style:25s}] [EMPTY + PAGE_BREAK]")
    # print all for structure
    elif style not in ['Normal', 'Default Paragraph Font', 'Normal (Web)']:
        print(f"[{idx:4d}] [{style:25s}] [empty]")

print("\n\n=== CHAPTER CONTEXT - Around chapter headings ===")
chapter_indices = [254, 256, 298, 334, 392, 441, 463, 518, 566, 590, 612, 622, 632]
for ch_idx in chapter_indices:
    print(f"\n--- Context around paragraph {ch_idx} ---")
    start = max(0, ch_idx - 3)
    end = min(len(doc.paragraphs), ch_idx + 5)
    for i in range(start, end):
        para = doc.paragraphs[i]
        text = para.text.strip()
        style = para.style.name
        pb = ""
        if 'pageBreak' in para._element.xml or 'w:type="page"' in para._element.xml:
            pb = " [PAGE_BREAK]"
        print(f"  [{i:4d}] [{style:25s}]{pb} |{text[:80]}|")
