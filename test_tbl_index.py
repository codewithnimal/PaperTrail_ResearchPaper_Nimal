import docx
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document(r'documentation.docs\Sample Project_Report.docx')
body = doc._body._element

# Let's locate Table 2, 3, 4
tbl2 = doc.tables[2]._element
tbl3 = doc.tables[3]._element
tbl4 = doc.tables[4]._element

print("Table 2 parent tag:", tbl2.getparent().tag)
print("Table 2 index in body:", body.index(tbl2))
print("Table 3 index in body:", body.index(tbl3))
print("Table 4 index in body:", body.index(tbl4))
