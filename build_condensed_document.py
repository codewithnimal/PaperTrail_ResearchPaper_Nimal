# build_condensed_document.py
# Generator for the Condensed Academic Project Documentation (Target: 35-40 pages)
# Project: AN AI-POWERED RESEARCH PAPER COMPARISON AND EVIDENCE VERIFICATION SYSTEM USING RETRIEVAL-AUGMENTED GENERATION (RAG)

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from report_helpers import set_cell_border, set_cell_margins, set_cell_shading
from chapters_data_condensed import (
    get_condensed_chapter_1, get_condensed_chapter_2, get_condensed_chapter_3,
    get_condensed_chapter_4, get_condensed_chapter_5, get_condensed_chapter_6,
    get_condensed_chapter_7, get_condensed_chapter_8, get_condensed_chapter_9,
    get_condensed_chapter_10, get_condensed_chapter_11, get_condensed_chapter_12
)
from chapters_data_part3 import get_references
from tables_data import get_tables_dict

TEMPLATE_PATH = r'documentation.docs\Sample Project_Report.docx'
OUTPUT_PATH = r'documentation.docs\PaperTrail_Project_Documentation_Condensed.docx'
PDF_OUTPUT_PATH = r'documentation.docs\PaperTrail_Project_Documentation_Condensed.pdf'

OFFICIAL_TITLE_L1 = "AN AI-POWERED RESEARCH PAPER COMPARISON AND"
OFFICIAL_TITLE_L2 = "EVIDENCE VERIFICATION SYSTEM USING"
OFFICIAL_TITLE_L3 = "RETRIEVAL-AUGMENTED GENERATION (RAG)"
OFFICIAL_TITLE_FULL = f"{OFFICIAL_TITLE_L1} {OFFICIAL_TITLE_L2} {OFFICIAL_TITLE_L3}"

RIGHT_MARGIN_TWIPS = 9000  # 6.25 inches printable area * 1440 twips

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=Pt(2), space_after=Pt(2),
          line_spacing=1.3, bold=False, italic=False, font_size=Pt(12), font_name='Times New Roman'):
    p = doc.add_paragraph()
    p.alignment = align
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(int(space_before.pt * 20)))
    sp.set(qn('w:after'), str(int(space_after.pt * 20)))
    sp.set(qn('w:line'), str(int(line_spacing * 240)))
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)

    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = font_size
    return p

def add_chapter_heading(doc, chapter_num, title):
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_num._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '240')  # 12 pt
    sp.set(qn('w:after'), '80')   # 4 pt
    pPr.append(sp)
    r1 = p_num.add_run(f"CHAPTER {chapter_num}")
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p_title._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '60')   # 3 pt
    sp2.set(qn('w:after'), '200')  # 10 pt
    pPr2.append(sp2)
    r2 = p_title.add_run(title)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '160')  # 8 pt
    sp.set(qn('w:after'), '60')   # 3 pt
    pPr.append(sp)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    return p

def add_code_block(doc, code_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '80')
    sp.set(qn('w:after'), '80')
    sp.set(qn('w:line'), '220')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '300')
    pPr.append(ind)
    r = p.add_run(code_str)
    r.font.name = 'Courier New'
    r.font.size = Pt(8.5)
    return p

def add_figure(doc, image_path, fig_no, caption, width=Inches(4.8)):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_img._element.get_or_add_pPr()
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '100')
        sp.set(qn('w:after'), '40')
        pPr.append(sp)
        run = p_img.add_run()
        run.add_picture(image_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p_cap._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '40')
    sp2.set(qn('w:after'), '120')
    pPr2.append(sp2)
    r_cap = p_cap.add_run(f"[{fig_no}: {caption}]")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10.5)
    r_cap.bold = False

def add_table(doc, tbl_data):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_cap._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '120')
    sp.set(qn('w:after'), '60')
    pPr.append(sp)
    r = p_cap.add_run(f"{tbl_data['num']}: {tbl_data['title']}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)

    headers = tbl_data['headers']
    rows = tbl_data['rows']
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    border_kwargs = dict(
        top=dict(sz=4, val='single', color='888888', space='0'),
        bottom=dict(sz=4, val='single', color='888888', space='0'),
        left=dict(sz=4, val='single', color='888888', space='0'),
        right=dict(sz=4, val='single', color='888888', space='0'),
        insideH=dict(sz=4, val='single', color='CCCCCC', space='0'),
        insideV=dict(sz=4, val='single', color='CCCCCC', space='0')
    )

    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.text = h_text
        set_cell_shading(cell, "EAEAEA")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_border(cell, **border_kwargs)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(8.5)

    for r_idx, row in enumerate(rows):
        for col_idx, val in enumerate(row):
            cell = tbl.cell(r_idx+1, col_idx)
            cell.text = str(val)
            if r_idx % 2 == 1:
                set_cell_shading(cell, "F9F9F9")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            set_cell_border(cell, **border_kwargs)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 and len(str(val)) > 20 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(8.5)

    p_sp = doc.add_paragraph()
    p_spPr = p_sp._element.get_or_add_pPr()
    sp_post = OxmlElement('w:spacing')
    sp_post.set(qn('w:before'), '0')
    sp_post.set(qn('w:after'), '100')
    p_spPr.append(sp_post)

def make_entry_paragraph(prefix, title, page_str, is_chapter=False, is_sub=False, pos=RIGHT_MARGIN_TWIPS):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')

    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '40' if is_chapter else '15')
    sp.set(qn('w:after'), '40' if is_chapter else '15')
    sp.set(qn('w:line'), '240')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)

    if is_sub:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '360')
        pPr.append(ind)

    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), str(pos))
    tabs.append(tab)
    pPr.append(tabs)

    p.append(pPr)

    r1 = OxmlElement('w:r')
    r1Pr = OxmlElement('w:rPr')
    rf1 = OxmlElement('w:rFonts')
    rf1.set(qn('w:ascii'), 'Times New Roman')
    rf1.set(qn('w:hAnsi'), 'Times New Roman')
    r1Pr.append(rf1)
    sz1 = OxmlElement('w:sz')
    sz1.set(qn('w:val'), '22' if is_chapter else '20')
    r1Pr.append(sz1)
    if is_chapter:
        r1Pr.append(OxmlElement('w:b'))
    r1.append(r1Pr)

    t1 = OxmlElement('w:t')
    label = f"{prefix}   {title}" if prefix else title
    t1.text = f"{label}\t"
    t1.set(qn('xml:space'), 'preserve')
    r1.append(t1)
    p.append(r1)

    r2 = OxmlElement('w:r')
    r2Pr = OxmlElement('w:rPr')
    rf2 = OxmlElement('w:rFonts')
    rf2.set(qn('w:ascii'), 'Times New Roman')
    rf2.set(qn('w:hAnsi'), 'Times New Roman')
    r2Pr.append(rf2)
    sz2 = OxmlElement('w:sz')
    sz2.set(qn('w:val'), '22' if is_chapter else '20')
    r2Pr.append(sz2)
    if is_chapter:
        r2Pr.append(OxmlElement('w:b'))
    r2.append(r2Pr)

    t2 = OxmlElement('w:t')
    t2.text = str(page_str)
    r2.append(t2)
    p.append(r2)

    return p

def update_front_matter(doc, toc_entries, lot_entries, lof_entries):
    print("Normalizing and updating front matter...")

    # Normalize all sections
    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.25)
        sec.right_margin = Inches(1.0)
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.5)

    # Title Page: Lines 1, 2, 3
    p1 = doc.paragraphs[1]
    p1.text = ""
    r1 = p1.add_run(OFFICIAL_TITLE_L1)
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(17)

    p2 = doc.paragraphs[2]
    p2.text = ""
    r2 = p2.add_run(OFFICIAL_TITLE_L2)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(17)

    p3 = doc.paragraphs[3]
    p3.text = ""
    r3 = p3.add_run(OFFICIAL_TITLE_L3)
    r3.bold = True
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(17)

    doc.paragraphs[4].text = ""
    doc.paragraphs[5].text = ""

    # Bonafide Certificate
    p27 = doc.paragraphs[27]
    p27.text = ""
    r27_1 = p27.add_run("Certified that this project report \"")
    r27_1.font.name = 'Times New Roman'
    r27_1.font.size = Pt(13)
    r27_2 = p27.add_run(OFFICIAL_TITLE_FULL)
    r27_2.bold = True
    r27_2.font.name = 'Times New Roman'
    r27_2.font.size = Pt(13)
    r27_3 = p27.add_run("\" is the bonafide work of SUBASH KUMAR K (723722243087), SIVAPRATHEESH S (723722243081), MOHAN KUMAR S (723722243302), and SURESH (723722243088) who carried out the project work under my supervision.")
    r27_3.font.name = 'Times New Roman'
    r27_3.font.size = Pt(13)

    # Acknowledgement: Ensure it fits cleanly on page iv by setting tight paragraph spacing
    doc.paragraphs[48].text = "ACKNOWLEDGEMENT"
    ack_paras = [49, 51, 53, 55, 59, 61]
    for p_idx in ack_paras:
        p = doc.paragraphs[p_idx]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '40')
        sp.set(qn('w:after'), '40')
        sp.set(qn('w:line'), '240')

    doc.paragraphs[49].text = "First, we extend our heartfelt gratitude to the management of V.S.B. College of Engineering Technical Campus for providing us with state-of-the-art computational infrastructure and research facilities."
    doc.paragraphs[51].text = "We express sincere thanks to our Principal, Dr. V. Velmurugan, M.E., Ph.D., for his continuous encouragement and academic leadership."
    doc.paragraphs[53].text = "We also thank our Vice Principal, Dr. R. Poornachandran, M.E., Ph.D., for his administrative cooperation and motivation."
    doc.paragraphs[55].text = "We are grateful to Dr. R. Murugadoss, M.E., Ph.D., and Mrs. G. Nithya, M.E., (Ph.D), Head of the Department of Artificial Intelligence and Data Science, for constructive suggestions and technical seminars."
    doc.paragraphs[59].text = "We express deepest gratitude to our project guide, Mr. S. Soundhar, M.E., Assistant Professor, Department of Artificial Intelligence and Data Science, for his invaluable technical guidance and meticulous architectural reviews throughout this project."
    doc.paragraphs[61].text = "Finally, we thank our parents, faculty members, and friends who supported us morally and technically in completing this research."

    # Abstract & Keywords: Tightened so it fits on page xi without spilling
    p159 = doc.paragraphs[159]
    p159.text = OFFICIAL_TITLE_FULL
    for r in p159.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    p172 = doc.paragraphs[172]
    p172.text = (
        "The rapid expansion of peer-reviewed scientific literature poses a formidable cognitive barrier for researchers, "
        "scholars, and peer reviewers. Traditional academic search engines index keywords but provide zero cross-document "
        "relational synthesis. Conversely, conversational Large Language Models (LLMs) frequently hallucinate citations, "
        "phantom benchmarks, and ungrounded claims. This project presents an enterprise-grade AI-Powered Research Paper Comparison "
        "and Evidence Verification System using Retrieval-Augmented Generation (RAG). The framework integrates an automated PDF "
        "ingestion pipeline using PyMuPDF preserving exact 1-indexed page coordinates, a metadata-preserving sliding window chunker "
        "(512 tokens with 100-token overlap), and a hybrid retrieval engine combining dense vector embeddings (SentenceTransformers "
        "all-MiniLM-L6-v2) with sparse BM25 lexical indexing fused via Reciprocal Rank Fusion (k=60). Claims are verified through "
        "Natural Language Inference (NLI) entailment filtering, generating clickable page citations formatted as [Paper Title, p. X]. "
        "The system constructs multi-dimensional comparative matrices spanning objectives, methodologies, architectures, datasets, "
        "metrics, findings, limitations, and research gaps, while an automated metric auditing engine verifies evaluation splits "
        "and measurement units. A Spring Security 6 module enforcing JWT and BCrypt hashing guarantees strict multi-tenant data isolation. "
        "Experimental evaluation across 50 scientific papers demonstrated a Top-5 retrieval precision of 96.8%, a page citation accuracy of 97.4%, "
        "and latency under 8 seconds for a 3-paper comparison, achieving 100% compliance across security test suites."
    )
    pPr172 = p172._element.get_or_add_pPr()
    sp172 = pPr172.find(qn('w:spacing'))
    if sp172 is None:
        sp172 = OxmlElement('w:spacing')
        pPr172.append(sp172)
    sp172.set(qn('w:before'), '60')
    sp172.set(qn('w:after'), '60')
    sp172.set(qn('w:line'), '240')  # 1.0 line spacing to guarantee single page

    for r in p172.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    p174 = doc.paragraphs[174]
    p174.text = (
        "Retrieval-Augmented Generation (RAG), scientific document intelligence, multi-paper comparison matrix, "
        "evidence grounding, page-level citations, metric consistency auditing, Spring Boot, Spring Security 6, JWT, PyMuPDF, FastAPI."
    )
    for r in p174.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    # Rebuild TOC, LOT, LOF
    body = doc._body._element

    # Ensure LIST OF TABLES starts on its own page
    for p in doc.paragraphs:
        if p.text.strip() == "LIST OF TABLES":
            p_break = OxmlElement('w:p')
            r_break = OxmlElement('w:r')
            br = OxmlElement('w:br')
            br.set(qn('w:type'), 'page')
            r_break.append(br)
            p_break.append(r_break)
            p._element.addprevious(p_break)
            break

    # Replace Table 2 (TOC)
    tbl2 = doc.tables[2]._element
    for item in toc_entries:
        p_el = make_entry_paragraph(item[0], item[1], item[2], item[3], item[4])
        tbl2.addprevious(p_el)
    body.remove(tbl2)

    # Replace Table 3 (LOT)
    tbl3 = doc.tables[2]._element
    for item in lot_entries:
        p_el = make_entry_paragraph(item[0], item[1], item[2], False, False)
        tbl3.addprevious(p_el)
    body.remove(tbl3)

    # Replace Table 4 (LOF)
    tbl4 = doc.tables[2]._element
    for item in lof_entries:
        p_el = make_entry_paragraph(item[0], item[1], item[2], False, False)
        tbl4.addprevious(p_el)
    body.remove(tbl4)

    print("Front matter updated successfully.")

def build_condensed_doc(toc_entries, lot_entries, lof_entries):
    print(f"Loading template: {TEMPLATE_PATH}...")
    doc = docx.Document(TEMPLATE_PATH)

    body = doc._body._element

    # In Sample Project_Report.docx:
    # Paragraph 184 has sectPr (end of Section 3, front matter list of figures)
    # Paragraphs 185 to 207 is the old separator ("CHAPTER - 1 \n INTRODUCTION") with sectPr (end of Section 4)
    # Paragraph 208 was where Section 5 began
    # We want to delete EVERYTHING from paragraph 185 up to the end of body, but preserve the final sectPr!
    ch_sep_elem = doc.paragraphs[185]._element

    # Update front matter first
    update_front_matter(doc, toc_entries, lot_entries, lof_entries)

    # Find where ch_sep_elem is in body now
    sep_idx = body.index(ch_sep_elem)
    print(f"Clearing old chapters and separators from body index {sep_idx} to end...")
    del body[sep_idx:-1]

    # Ensure the final sectPr sets page start="1" with decimal format
    final_sectPr = body[-1]
    pgNumType = final_sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        final_sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), '1')
    pgNumType.set(qn('w:fmt'), 'decimal')

    tables_dict = get_tables_dict()

    chapters = [
        get_condensed_chapter_1(),
        get_condensed_chapter_2(),
        get_condensed_chapter_3(),
        get_condensed_chapter_4(),
        get_condensed_chapter_5(),
        get_condensed_chapter_6(),
        get_condensed_chapter_7(),
        get_condensed_chapter_8(),
        get_condensed_chapter_9(),
        get_condensed_chapter_10(),
        get_condensed_chapter_11(),
        get_condensed_chapter_12()
    ]

    print("Appending condensed chapters, tables, and figures...")

    for ch in chapters:
        ch_num = ch["num"]
        ch_title = ch["title"]
        print(f"Writing Chapter {ch_num}: {ch_title}...")
        add_chapter_heading(doc, ch_num, ch_title)

        for sec in ch["sections"]:
            sec_num = sec["num"]
            sec_title = sec["title"]
            add_h2(doc, f"{sec_num}   {sec_title}")

            for p_text in sec["paragraphs"]:
                if p_text.startswith("Algorithm 4."):
                    add_code_block(doc, p_text)
                else:
                    add_p(doc, p_text)

            # Insert Figures and Tables at their academic locations
            if sec_num == "2.5":
                add_table(doc, tables_dict["table_2_1"])
            elif sec_num == "3.1":
                add_figure(doc, "generated_figures/fig_3_1_architecture.png", "Figure 3.1", "End-to-End System Architecture: AI-Powered Research Paper Comparison Engine", width=Inches(4.8))
            elif sec_num == "3.2":
                add_figure(doc, "generated_figures/fig_3_2_pdf_extraction.png", "Figure 3.2", "PDF Ingestion and Page-Bound Text Extraction Pipeline", width=Inches(4.8))
            elif sec_num == "3.3":
                add_figure(doc, "generated_figures/fig_3_3_chunking.png", "Figure 3.3", "Metadata-Preserving Hierarchical Chunking with Sliding Window Overlap", width=Inches(4.8))
            elif sec_num == "3.4":
                add_figure(doc, "generated_figures/fig_3_4_retrieval.png", "Figure 3.4", "Hybrid Semantic-Lexical Retrieval Architecture with Reciprocal Rank Fusion", width=Inches(4.8))
            elif sec_num == "3.5":
                add_figure(doc, "generated_figures/fig_3_5_auth_flow.png", "Figure 3.5", "User Authentication and JWT Authorization Flow with Spring Security", width=Inches(4.8))
            elif sec_num == "3.6":
                add_figure(doc, "generated_figures/fig_3_6_data_isolation.png", "Figure 3.6", "Multi-Tenant Data Isolation and Ownership Verification Architecture", width=Inches(4.8))
            elif sec_num == "4.1":
                add_table(doc, tables_dict["table_4_1"])
            elif sec_num == "5.1":
                add_table(doc, tables_dict["table_5_1"])
            elif sec_num == "5.2":
                add_table(doc, tables_dict["table_5_2"])
            elif sec_num == "5.3":
                add_table(doc, tables_dict["table_5_3"])
            elif sec_num == "6.1":
                add_figure(doc, "generated_figures/fig_6_1_comparison_pipeline.png", "Figure 6.1", "Multi-Paper Comparison Matrix and Claim Verification Pipeline", width=Inches(4.8))
            elif sec_num == "6.3":
                add_table(doc, tables_dict["table_6_1"])
            elif sec_num == "6.5":
                add_figure(doc, "generated_figures/fig_6_2_metric_verification.png", "Figure 6.2", "Metric Extraction and Cross-Paper Consistency Verification Pipeline", width=Inches(4.8))
            elif sec_num == "6.6":
                add_figure(doc, "generated_figures/fig_6_3_async_jobs.png", "Figure 6.3", "Asynchronous Job Processing and State Transition Lifecycle", width=Inches(4.8))
            elif sec_num == "7.2":
                add_table(doc, tables_dict["table_7_1"])
            elif sec_num == "7.3":
                add_figure(doc, "generated_figures/fig_7_1_precision_benchmark.png", "Figure 7.1", "Evidence Retrieval Precision vs. Token Chunk Size Benchmark", width=Inches(4.6))
                add_table(doc, tables_dict["table_7_2"])
            elif sec_num == "7.4":
                add_figure(doc, "generated_figures/fig_7_2_latency_benchmark.png", "Figure 7.2", "Inference and Matrix Construction Latency vs. Number of Papers", width=Inches(4.6))
                add_table(doc, tables_dict["table_7_3"])
            elif sec_num == "7.5":
                add_table(doc, tables_dict["table_7_4"])
            elif sec_num == "7.6":
                add_figure(doc, "generated_figures/fig_7_3_entailment_heatmap.png", "Figure 7.3", "Cross-Paper Claim Grounding Entailment Scores by Dimension Heatmap", width=Inches(4.6))

    print("Writing References...")
    p_ref_h = doc.add_paragraph()
    p_ref_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ref_hPr = p_ref_h._element.get_or_add_pPr()
    sp_ref = OxmlElement('w:spacing')
    sp_ref.set(qn('w:before'), '240')
    sp_ref.set(qn('w:after'), '140')
    p_ref_hPr.append(sp_ref)
    r_ref = p_ref_h.add_run("REFERENCES")
    r_ref.bold = True
    r_ref.font.name = 'Times New Roman'
    r_ref.font.size = Pt(15)

    references = get_references()
    for ref_str in references:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pPr = p_ref._element.get_or_add_pPr()
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '20')
        sp.set(qn('w:after'), '30')
        sp.set(qn('w:line'), '240')
        sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '400')
        ind.set(qn('w:hanging'), '400')
        pPr.append(ind)
        r = p_ref.add_run(ref_str)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)

    print(f"Saving to {OUTPUT_PATH}...")
    doc.save(OUTPUT_PATH)
    print("Saved successfully.")

if __name__ == '__main__':
    # Accurate page numbers derived from PDF inspection
    dummy_toc = [
        ("", "LIST OF TABLES", "xii", True, False),
        ("", "LIST OF FIGURES", "xiii", True, False),
        ("CHAPTER 1", "INTRODUCTION", "1", True, False),
        ("1.1", "Background and Context", "1", False, True),
        ("1.2", "Problem Statement", "1", False, True),
        ("1.3", "Objectives of the Project", "2", False, True),
        ("1.4", "Scope of the Project", "3", False, True),
        ("1.5", "Authentication and User Data Isolation Requirements", "3", False, True),
        ("CHAPTER 2", "LITERATURE REVIEW", "3", True, False),
        ("2.1", "Evolution of Scientific Document Processing", "3", False, True),
        ("2.2", "Traditional NLP and Extractive Summarization", "3", False, True),
        ("2.3", "Large Language Models and Hallucination in Academic QA", "4", False, True),
        ("2.4", "Retrieval-Augmented Generation (RAG) in Scientific Domains", "4", False, True),
        ("2.5", "Comparative Analysis of Existing Literature Review Assistants", "4", False, True),
        ("2.6", "Summary of Research Gaps", "5", False, True),
        ("CHAPTER 3", "SYSTEM DESIGN AND ARCHITECTURE", "5", True, False),
        ("3.1", "Overall System Architecture", "5", False, True),
        ("3.2", "PDF Ingestion and Page-Bound Text Extraction Pipeline", "6", False, True),
        ("3.3", "Metadata-Preserving Hierarchical Chunking", "6", False, True),
        ("3.4", "Hybrid Semantic-Lexical Retrieval Architecture", "7", False, True),
        ("3.5", "Authentication, Authorization, and Security Architecture", "7", False, True),
        ("3.6", "Multi-Tenant Data Isolation and Ownership Verification Model", "8", False, True),
        ("CHAPTER 4", "RESEARCH METHODOLOGY AND ALGORITHMIC FORMULATION", "8", True, False),
        ("CHAPTER 5", "SYSTEM REQUIREMENTS", "11", True, False),
        ("CHAPTER 6", "SYSTEM DESIGN AND IMPLEMENTATION", "12", True, False),
        ("CHAPTER 7", "EXPERIMENTAL RESULTS AND ANALYSIS", "15", True, False),
        ("CHAPTER 8", "USE CASES", "18", True, False),
        ("CHAPTER 9", "ETHICAL, LEGAL AND SOCIETAL CONSIDERATIONS", "19", True, False),
        ("CHAPTER 10", "LIMITATIONS", "20", True, False),
        ("CHAPTER 11", "FUTURE SCOPE", "20", True, False),
        ("CHAPTER 12", "CONCLUSION", "21", True, False),
        ("", "REFERENCES", "21", True, False)
    ]
    dummy_lot = [
        ("Table 2.1", "Comparative Analysis of Literature Review Tools and Proposed System", "4"),
        ("Table 4.1", "Benchmark Evaluation Dataset and Scientific Document Corpus", "9"),
        ("Table 5.1", "Hardware Requirements", "11"),
        ("Table 5.2", "Software Requirements and Technology Stack", "11"),
        ("Table 5.3", "Security, Authentication, and Compliance Requirements Specification", "12"),
        ("Table 6.1", "Key Backend REST API Endpoints and Authorization Rules", "13"),
        ("Table 7.1", "Evidence Retrieval Precision and Page Grounding Accuracy", "15"),
        ("Table 7.2", "Ablation Study Results on Chunking Parameters and Fusion", "16"),
        ("Table 7.3", "Inference Latency Breakdown for a 3-Paper Comparison Session", "16"),
        ("Table 7.4", "Authentication, Authorization, and Security Test Suite Results", "17")
    ]
    dummy_lof = [
        ("Fig. 3.1", "End-to-End System Architecture: AI-Powered Research Paper Comparison Engine", "5"),
        ("Fig. 3.2", "PDF Ingestion and Page-Bound Text Extraction Pipeline", "6"),
        ("Fig. 3.3", "Metadata-Preserving Hierarchical Chunking with Sliding Window Overlap", "6"),
        ("Fig. 3.4", "Hybrid Semantic-Lexical Retrieval Architecture with Reciprocal Rank Fusion", "7"),
        ("Fig. 3.5", "User Authentication and JWT Authorization Flow with Spring Security", "7"),
        ("Fig. 3.6", "Multi-Tenant Data Isolation and Ownership Verification Architecture", "8"),
        ("Fig. 6.1", "Multi-Paper Comparison Matrix and Claim Verification Pipeline", "13"),
        ("Fig. 6.2", "Metric Extraction and Cross-Paper Consistency Verification Pipeline", "14"),
        ("Fig. 6.3", "Asynchronous Job Processing and State Transition Lifecycle", "15"),
        ("Fig. 7.1", "Evidence Retrieval Precision vs. Token Chunk Size Benchmark", "16"),
        ("Fig. 7.2", "Inference and Matrix Construction Latency vs. Number of Papers", "16"),
        ("Fig. 7.3", "Cross-Paper Claim Grounding Entailment Scores by Dimension Heatmap", "18")
    ]
    build_condensed_doc(dummy_toc, dummy_lot, dummy_lof)
