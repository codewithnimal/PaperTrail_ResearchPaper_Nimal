import docx

doc = docx.Document(r'documentation.docs\Sample Project_Report.docx')
body = doc._body._element

ch1_p = doc.paragraphs[206]
print("P 206 text:", repr(ch1_p.text))
print("Index in body:", body.index(ch1_p._element))

ch1_p2 = doc.paragraphs[208]
print("P 208 text:", repr(ch1_p2.text))
print("Index in body:", body.index(ch1_p2._element))
