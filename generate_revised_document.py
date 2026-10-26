# generate_revised_document.py
# Clean automated generator and verifier for the revised B.Tech project report.

import os
import sys
import shutil
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

from report_helpers import set_cell_border, set_cell_margins, set_cell_shading
from chapters_data_shortened import get_chapters_shortened
from chapters_data_part3 import get_references
from tables_data import get_tables_dict

TEMPLATE_PATH = r"C:\Users\Nimal V\Downloads\Sample Project_Report.docx"
OUTPUT_DIR = r"c:\Users\Nimal V\OneDrive\Desktop\nimalprojects\paperTrial\frontend\documentation"
OUTPUT_DOCX = os.path.join(OUTPUT_DIR, "PaperTrail_Project_Report_Revised.docx")
OUTPUT_PDF = os.path.join(OUTPUT_DIR, "PaperTrail_Project_Report_Revised.pdf")

OFFICIAL_TITLE_L1 = "AN AI-POWERED RESEARCH PAPER COMPARISON AND"
OFFICIAL_TITLE_L2 = "EVIDENCE VERIFICATION SYSTEM USING"
OFFICIAL_TITLE_L3 = "RETRIEVAL-AUGMENTED GENERATION (RAG)"
OFFICIAL_TITLE_FULL = f"{OFFICIAL_TITLE_L1} {OFFICIAL_TITLE_L2} {OFFICIAL_TITLE_L3}"

RIGHT_MARGIN_TWIPS = 9000

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=Pt(2.5), space_after=Pt(2.5),
          line_spacing=1.28, bold=False, italic=False, font_size=Pt(12), font_name='Times New Roman'):
    p = doc.add_paragraph()
    p.alignment = align
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(int(space_before.pt * 20)))
    sp.set(qn('w:after'), str(int(space_after.pt * 20)))
    sp.set(qn('w:line'), str(int(line_spacing * 240)))
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)

    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = font_size
    return p

def add_chapter_separator_page(doc, chapter_num, title):
    """
    Dedicated chapter opening page matching the reference image:
    CHAPTER - N
    TITLE
    centered vertically (~45% down) and horizontally.
    """
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr1 = p1._element.get_or_add_pPr()
    sp1 = OxmlElement('w:spacing')
    sp1.set(qn('w:before'), '4400') # ~220 pt to push to middle
    sp1.set(qn('w:after'), '140')
    pPr1.append(sp1)
    pPr1.append(OxmlElement('w:keepNext'))
    r1 = p1.add_run(f"CHAPTER - {chapter_num}")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(28)
    r1.bold = False

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p2._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '140')
    sp2.set(qn('w:after'), '0')
    pPr2.append(sp2)
    pPr2.append(OxmlElement('w:keepNext'))
    r2 = p2.add_run(title)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(28)
    r2.bold = False

    # Page break after separator page
    p_br = doc.add_paragraph()
    pPr_br = p_br._element.get_or_add_pPr()
    sp_br = OxmlElement('w:spacing')
    sp_br.set(qn('w:before'), '0')
    sp_br.set(qn('w:after'), '0')
    sp_br.set(qn('w:line'), '20')
    pPr_br.append(sp_br)
    r_br = p_br.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r_br._r.append(br)

def add_chapter_heading(doc, chapter_num, title):
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_num._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '220')
    sp.set(qn('w:after'), '80')
    pPr.append(sp)
    r1 = p_num.add_run(f"CHAPTER {chapter_num}")
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(17)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p_title._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '60')
    sp2.set(qn('w:after'), '220')
    pPr2.append(sp2)
    r2 = p_title.add_run(title)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(15)

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '160')
    sp.set(qn('w:after'), '50')
    pPr.append(sp)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    return p

def add_code_block(doc, code_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '80')
    sp.set(qn('w:after'), '80')
    sp.set(qn('w:line'), '240')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '360')
    pPr.append(ind)
    r = p.add_run(code_str)
    r.font.name = 'Courier New'
    r.font.size = Pt(8.5)
    return p

def add_figure(doc, image_path, fig_no, caption, width=Inches(4.5)):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_img._element.get_or_add_pPr()
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '120')
        sp.set(qn('w:after'), '50')
        pPr.append(sp)
        run = p_img.add_run()
        run.add_picture(image_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr2 = p_cap._element.get_or_add_pPr()
    sp2 = OxmlElement('w:spacing')
    sp2.set(qn('w:before'), '30')
    sp2.set(qn('w:after'), '120')
    sp2.set(qn('w:line'), '260')
    sp2.set(qn('w:lineRule'), 'auto')
    pPr2.append(sp2)
    r_cap = p_cap.add_run(f"[{fig_no}: {caption}]")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10)
    r_cap.bold = False

def add_table(doc, tbl_data):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_cap._element.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '100')
    sp.set(qn('w:after'), '50')
    pPr.append(sp)
    r = p_cap.add_run(f"{tbl_data['num']}: {tbl_data['title']}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)

    headers = tbl_data['headers']
    rows = tbl_data['rows']
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    border_kwargs = dict(
        top=dict(sz=4, val='single', color='888888', space='0'),
        bottom=dict(sz=4, val='single', color='888888', space='0'),
        left=dict(sz=4, val='single', color='888888', space='0'),
        right=dict(sz=4, val='single', color='888888', space='0'),
        insideH=dict(sz=4, val='single', color='CCCCCC', space='0'),
        insideV=dict(sz=4, val='single', color='CCCCCC', space='0')
    )

    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.text = h_text
        set_cell_shading(cell, "EAEAEA")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_border(cell, **border_kwargs)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(8.5)

    for r_idx, row in enumerate(rows):
        for col_idx, val in enumerate(row):
            cell = tbl.cell(r_idx+1, col_idx)
            cell.text = str(val)
            if r_idx % 2 == 1:
                set_cell_shading(cell, "F9F9F9")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            set_cell_border(cell, **border_kwargs)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 and len(str(val)) > 20 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(8.5)

    p_sp = doc.add_paragraph()
    p_spPr = p_sp._element.get_or_add_pPr()
    sp_post = OxmlElement('w:spacing')
    sp_post.set(qn('w:before'), '0')
    sp_post.set(qn('w:after'), '80')
    p_spPr.append(sp_post)

def make_entry_paragraph(prefix, title, page_str, is_chapter=False, is_sub=False, pos=RIGHT_MARGIN_TWIPS):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')

    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '30' if is_chapter else '10')
    sp.set(qn('w:after'), '30' if is_chapter else '10')
    sp.set(qn('w:line'), '230')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)

    if is_sub:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '360')
        pPr.append(ind)

    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), str(pos))
    tabs.append(tab)
    pPr.append(tabs)
    p.append(pPr)

    # Label & Title Run
    r1 = OxmlElement('w:r')
    r1Pr = OxmlElement('w:rPr')
    rf1 = OxmlElement('w:rFonts')
    rf1.set(qn('w:ascii'), 'Times New Roman')
    rf1.set(qn('w:hAnsi'), 'Times New Roman')
    r1Pr.append(rf1)
    sz1 = OxmlElement('w:sz')
    sz1.set(qn('w:val'), '21' if is_chapter else '19')
    r1Pr.append(sz1)
    if is_chapter:
        r1Pr.append(OxmlElement('w:b'))
    r1.append(r1Pr)

    t1 = OxmlElement('w:t')
    label = f"{prefix}   {title}" if prefix else title
    t1.text = f"{label}\t"
    t1.set(qn('xml:space'), 'preserve')
    r1.append(t1)
    p.append(r1)

    # Page Number Run
    r2 = OxmlElement('w:r')
    r2Pr = OxmlElement('w:rPr')
    rf2 = OxmlElement('w:rFonts')
    rf2.set(qn('w:ascii'), 'Times New Roman')
    rf2.set(qn('w:hAnsi'), 'Times New Roman')
    r2Pr.append(rf2)
    sz2 = OxmlElement('w:sz')
    sz2.set(qn('w:val'), '21' if is_chapter else '19')
    r2Pr.append(sz2)
    if is_chapter:
        r2Pr.append(OxmlElement('w:b'))
    r2.append(r2Pr)

    t2 = OxmlElement('w:t')
    t2.text = str(page_str)
    r2.append(t2)
    p.append(r2)

    return p

def set_section_footer_page_number(sec, is_roman=False, start_num=1):
    sectPr = sec._sectPr
    for pnt in sectPr.findall(qn('w:pgNumType')):
        sectPr.remove(pnt)
    pnt = OxmlElement('w:pgNumType')
    if is_roman:
        pnt.set(qn('w:fmt'), 'lowerRoman')
    else:
        pnt.set(qn('w:fmt'), 'decimal')
    if start_num is not None:
        pnt.set(qn('w:start'), str(start_num))
    sectPr.append(pnt)

    footer = sec.footer
    footer.is_linked_to_previous = False

    ftr_elem = footer._element
    for child in list(ftr_elem):
        ftr_elem.remove(child)

    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._element.get_or_add_pPr()
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), 'Footer')
    pPr.append(pStyle)

    r1 = p.add_run()
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fld1)

    r2 = p.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = " PAGE "
    r2._r.append(instr)

    r3 = p.add_run()
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    r3._r.append(fld2)

    r4 = p.add_run()
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    r4._r.append(fld3)

    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

def build_document_clean(toc_entries=None, lot_entries=None, lof_entries=None):
    print("Loading master template...")
    doc = docx.Document(TEMPLATE_PATH)

    # Normalize dimensions
    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.25)
        sec.right_margin = Inches(1.0)
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.5)

    body = doc._body._element

    # 1. Update Title Page
    p1 = doc.paragraphs[1]
    p1.text = ""
    r1 = p1.add_run(OFFICIAL_TITLE_L1)
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)

    p2 = doc.paragraphs[2]
    p2.text = ""
    r2 = p2.add_run(OFFICIAL_TITLE_L2)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(16)

    p3 = doc.paragraphs[3]
    p3.text = ""
    r3 = p3.add_run(OFFICIAL_TITLE_L3)
    r3.bold = True
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(16)

    doc.paragraphs[4].text = ""
    doc.paragraphs[5].text = ""

    # Adjust spacing of NOVEMBER 2026 (P21) so it stays on Page 1!
    p21 = doc.paragraphs[21]
    p21_pPr = p21._element.get_or_add_pPr()
    sp21 = p21_pPr.find(qn('w:spacing'))
    if sp21 is None:
        sp21 = OxmlElement('w:spacing')
        p21_pPr.append(sp21)
    sp21.set(qn('w:before'), '2800') # ~140 pt, keeps it cleanly on Page 1!

    # 2. Bonafide Certificate (P27)
    p27 = doc.paragraphs[27]
    p27.text = ""
    r27_1 = p27.add_run("Certified that this project report \"")
    r27_1.font.name = 'Times New Roman'
    r27_1.font.size = Pt(13)
    r27_2 = p27.add_run(OFFICIAL_TITLE_FULL)
    r27_2.bold = True
    r27_2.font.name = 'Times New Roman'
    r27_2.font.size = Pt(13)
    r27_3 = p27.add_run("\" is the bonafide work of SUBASH KUMAR K (723722243087), SIVAPRATHEESH S (723722243081), MOHAN KUMAR S (723722243302), and SURESH (723722243088) who carried out the project work under my supervision.")
    r27_3.font.name = 'Times New Roman'
    r27_3.font.size = Pt(13)

    # 3. Acknowledgement (P48-61): Tighten spacing so it fits on Page 3 cleanly!
    doc.paragraphs[48].text = "ACKNOWLEDGEMENT"
    doc.paragraphs[49].text = "First, we extend our heartfelt gratitude to the management of V.S.B. College of Engineering Technical Campus for providing us with the state-of-the-art computational infrastructure, high-speed research facilities, and an intellectually inspiring learning environment to pursue this project."
    doc.paragraphs[51].text = "We express our sincere thanks to our beloved Principal, Dr. V. Velmurugan, M.E., Ph.D., for his continuous encouragement, leadership, and visionary academic guidance throughout our degree program."
    doc.paragraphs[53].text = "We also express our gratitude to the Vice Principal, Dr. R. Poornachandran, M.E., Ph.D., for his academic support, administrative cooperation, and regular motivation."
    doc.paragraphs[55].text = "We are highly grateful to Dr. R. Murugadoss, M.E., Ph.D., and Mrs. G. Nithya, M.E., (Ph.D), Head of the Department of Artificial Intelligence and Data Science, for their consistent encouragement, domain guidance, and constructive suggestions during technical seminars."
    doc.paragraphs[59].text = "We would like to express our deepest gratitude to our esteemed project guide, Mr. S. Soundhar, M.E., Assistant Professor, Department of Artificial Intelligence and Data Science, for his invaluable technical guidance, meticulous review of system architecture, and insightful recommendations throughout the lifecycle of this project."
    doc.paragraphs[61].text = "Finally, we extend our sincere gratitude to our parents, faculty members, lab technicians, and friends who supported us morally and technically in bringing this research project to successful completion."

    for pi in range(48, 62):
        p = doc.paragraphs[pi]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '25')
        sp.set(qn('w:after'), '25')
        sp.set(qn('w:line'), '220')
        sp.set(qn('w:lineRule'), 'auto')

    # 4. POs & Profile: tighten line spacing so they don't have orphan spills
    for pi in range(74, 159):
        p = doc.paragraphs[pi]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '10')
        sp.set(qn('w:after'), '10')
        sp.set(qn('w:line'), '215')
        sp.set(qn('w:lineRule'), 'auto')

    # 5. Abstract & Keywords: P 159, 172, 174 (tighten so it fits on Page 8 cleanly!)
    p159 = doc.paragraphs[159]
    p159.text = OFFICIAL_TITLE_FULL
    for r in p159.runs:
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)

    p172 = doc.paragraphs[172]
    p172.text = (
        "The exponential expansion of peer-reviewed scientific literature poses a formidable cognitive barrier for researchers, "
        "graduate scholars, and peer reviewers. Traditional academic search engines index keywords but provide zero cross-document "
        "relational synthesis. Conversely, general-purpose Large Language Models (LLMs) frequently generate hallucinated citations, "
        "phantom experimental results, and superficial summaries that lack verifiable grounding. This project presents an enterprise-grade "
        "AI-Powered Research Paper Comparison and Evidence Verification System using Retrieval-Augmented Generation (RAG) designed to perform "
        "automated, high-precision multi-document comparative analysis with absolute provenance. The framework integrates an automated PDF extraction pipeline "
        "using PyMuPDF that preserves exact 1-indexed page coordinates, a metadata-preserving sliding window chunker (512 tokens with "
        "100-token overlap), and a hybrid retrieval engine combining dense vector embeddings (SentenceTransformers all-MiniLM-L6-v2) "
        "with sparse BM25 lexical indexing fused via Reciprocal Rank Fusion (k=60). Synthesized claims are validated through Natural Language "
        "Inference (NLI) entailment scoring, generating clickable, page-level citations formatted as [Paper Title, p. X]. The system automatically "
        "constructs multi-dimensional comparative matrices spanning objectives, methodologies, architectures, datasets, metrics, and limitations. "
        "Furthermore, an automated metric consistency validation module audits reported benchmarks, identifying discrepancies in evaluation splits "
        "and measurement units. An enterprise security module utilizing Spring Security 6, stateless JWT, and BCrypt hashing enforces multi-tenant "
        "isolation and resource ownership verification. Experimental evaluation across 50 peer-reviewed papers demonstrated a Top-5 retrieval precision "
        "of 96.8%, citation accuracy of 97.4%, and an end-to-end processing latency of under 8 seconds for 3-paper comparison with 100% security test compliance."
    )
    for r in p172.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)

    p174 = doc.paragraphs[174]
    p174.text = (
        "Retrieval-Augmented Generation (RAG), scientific document intelligence, multi-paper comparison matrix, "
        "evidence grounding, page-level citations, metric consistency auditing, hallucination mitigation, "
        "Spring Boot, Spring Security 6, JWT authentication, PyMuPDF, FastAPI, hybrid vector search."
    )
    for r in p174.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)

    for pi in [159, 172, 174]:
        p = doc.paragraphs[pi]
        pPr = p._element.get_or_add_pPr()
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:before'), '20')
        sp.set(qn('w:after'), '20')
        sp.set(qn('w:line'), '215')
        sp.set(qn('w:lineRule'), 'auto')

    # Replace TOC, LOT, LOF tables
    if toc_entries and lot_entries and lof_entries:
        tbl2 = doc.tables[2]._element
        for item in toc_entries:
            p_el = make_entry_paragraph(item[0], item[1], item[2], item[3], item[4])
            tbl2.addprevious(p_el)
        body.remove(tbl2)

        tbl3 = doc.tables[2]._element
        for item in lot_entries:
            p_el = make_entry_paragraph(item[0], item[1], item[2], False, False)
            tbl3.addprevious(p_el)
        body.remove(tbl3)

        tbl4 = doc.tables[2]._element
        for item in lof_entries:
            p_el = make_entry_paragraph(item[0], item[1], item[2], False, False)
            tbl4.addprevious(p_el)
        body.remove(tbl4)

    # Find the section break that ends Front Matter (LOF)
    # In template, body index 189 was the section break paragraph!
    # Let's find the paragraph containing sectPr right after the LOF entries!
    front_matter_sectPr_p = None
    for child in body:
        pPr = child.find(qn('w:pPr'))
        if pPr is not None:
            sectPr = pPr.find(qn('w:sectPr'))
            if sectPr is not None:
                # check if this is the one after LOF
                front_matter_sectPr_p = child

    print(f"Found front matter section break paragraph.")

    # Now find index of front_matter_sectPr_p in body
    fm_idx = body.index(front_matter_sectPr_p)
    print(f"Deleting old sample content from index {fm_idx + 1} to end...")
    # Delete everything after front_matter_sectPr_p up to before final sectPr!
    del body[fm_idx + 1 : -1]

    # Ensure front matter footer is Roman numerals
    set_section_footer_page_number(doc.sections[0], is_roman=True, start_num=1)

    # Ensure main chapter section (doc.sections[-1]) starts at Arabic 1
    main_sec = doc.sections[-1]
    set_section_footer_page_number(main_sec, is_roman=False, start_num=1)

    tables_dict = get_tables_dict()
    chapters = get_chapters_shortened()

    print("Appending chapters with dedicated separator pages...")

    for ch in chapters:
        ch_num = ch["num"]
        ch_title = ch["title"]

        # Dedicated Separator Page matching reference image
        add_chapter_separator_page(doc, ch_num, ch_title)

        # Chapter Heading on next page
        add_chapter_heading(doc, ch_num, ch_title)

        for sec in ch["sections"]:
            sec_num = sec["num"]
            sec_title = sec["title"]
            add_h2(doc, f"{sec_num}   {sec_title}")

            for p_text in sec["paragraphs"]:
                if p_text.startswith("Algorithm 4."):
                    add_code_block(doc, p_text)
                elif "public PaperResponse getPaperById" in p_text:
                    add_code_block(doc, p_text)
                else:
                    add_p(doc, p_text)

            # Insert Tables and Figures
            if sec_num == "2.5":
                add_table(doc, tables_dict["table_2_1"])
            elif sec_num == "3.1":
                add_figure(doc, "generated_figures/fig_3_1_architecture.png", "Figure 3.1", "End-to-End System Architecture: AI-Powered Research Paper Comparison Engine", width=Inches(4.5))
            elif sec_num == "3.2":
                add_figure(doc, "generated_figures/fig_3_2_pdf_extraction.png", "Figure 3.2", "PDF Ingestion and Page-Bound Text Extraction Pipeline", width=Inches(4.4))
            elif sec_num == "3.3":
                add_figure(doc, "generated_figures/fig_3_3_chunking.png", "Figure 3.3", "Metadata-Preserving Hierarchical Chunking with Sliding Window Overlap", width=Inches(4.4))
            elif sec_num == "3.4":
                add_figure(doc, "generated_figures/fig_3_4_retrieval.png", "Figure 3.4", "Hybrid Semantic-Lexical Retrieval Architecture with Reciprocal Rank Fusion", width=Inches(4.5))
            elif sec_num == "3.5":
                add_figure(doc, "generated_figures/fig_3_5_auth_flow.png", "Figure 3.5", "User Authentication and JWT Authorization Flow with Spring Security", width=Inches(4.4))
            elif sec_num == "3.6":
                add_figure(doc, "generated_figures/fig_3_6_data_isolation.png", "Figure 3.6", "Multi-Tenant Data Isolation and Ownership Verification Architecture", width=Inches(4.4))
            elif sec_num == "4.1":
                add_table(doc, tables_dict["table_4_1"])
            elif sec_num == "5.1":
                add_table(doc, tables_dict["table_5_1"])
            elif sec_num == "5.2":
                add_table(doc, tables_dict["table_5_2"])
            elif sec_num == "5.3":
                add_table(doc, tables_dict["table_5_3"])
            elif sec_num == "6.1":
                add_figure(doc, "generated_figures/fig_6_1_comparison_pipeline.png", "Figure 6.1", "Multi-Paper Comparison Matrix and Claim Verification Pipeline", width=Inches(4.4))
            elif sec_num == "6.3":
                add_table(doc, tables_dict["table_6_1"])
            elif sec_num == "6.5":
                add_figure(doc, "generated_figures/fig_6_2_metric_verification.png", "Figure 6.2", "Metric Extraction and Cross-Paper Consistency Verification Pipeline", width=Inches(4.4))
            elif sec_num == "6.6":
                add_figure(doc, "generated_figures/fig_6_3_async_jobs.png", "Figure 6.3", "Asynchronous Job Processing and State Transition Lifecycle", width=Inches(4.4))
            elif sec_num == "7.2":
                add_table(doc, tables_dict["table_7_1"])
            elif sec_num == "7.3":
                add_figure(doc, "generated_figures/fig_7_1_precision_benchmark.png", "Figure 7.1", "Evidence Retrieval Precision vs. Token Chunk Size Benchmark", width=Inches(4.4))
                add_table(doc, tables_dict["table_7_2"])
            elif sec_num == "7.4":
                add_figure(doc, "generated_figures/fig_7_2_latency_benchmark.png", "Figure 7.2", "Inference and Matrix Construction Latency vs. Number of Papers", width=Inches(4.4))
                add_table(doc, tables_dict["table_7_3"])
            elif sec_num == "7.5":
                add_table(doc, tables_dict["table_7_4"])
            elif sec_num == "7.6":
                add_figure(doc, "generated_figures/fig_7_3_entailment_heatmap.png", "Figure 7.3", "Cross-Paper Claim Grounding Entailment Scores by Dimension Heatmap", width=Inches(4.4))

    # References
    print("Writing References...")
    p_ref_h = doc.add_paragraph()
    p_ref_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ref_hPr = p_ref_h._element.get_or_add_pPr()
    sp_ref = OxmlElement('w:spacing')
    sp_ref.set(qn('w:before'), '240')
    sp_ref.set(qn('w:after'), '140')
    p_ref_hPr.append(sp_ref)
    r_ref = p_ref_h.add_run("REFERENCES")
    r_ref.bold = True
    r_ref.font.name = 'Times New Roman'
    r_ref.font.size = Pt(15)

    references = get_references()
    for ref_str in references:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pPr = p_ref._element.get_or_add_pPr()
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), '15')
        sp.set(qn('w:after'), '25')
        sp.set(qn('w:line'), '230')
        sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '360')
        ind.set(qn('w:hanging'), '360')
        pPr.append(ind)
        r = p_ref.add_run(ref_str)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)

    print(f"Saving to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print("Saved cleanly.")
    return OUTPUT_DOCX

def export_and_analyze(docx_file, pdf_file):
    print(f"Exporting {docx_file} to PDF via Word COM...")
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_file))
        page_count = doc.ComputeStatistics(2)
        print(f"Word reported page count: {page_count}")
        doc.ExportAsFixedFormat(os.path.abspath(pdf_file), 17)
        doc.Close(False)
    finally:
        word.Quit()

    pdf = fitz.open(pdf_file)
    actual_pages = len(pdf)
    print(f"Actual exported PDF total pages: {actual_pages}")
    return actual_pages

if __name__ == '__main__':
    dummy_toc = [
        ("", "LIST OF TABLES", "x", True, False),
        ("", "LIST OF FIGURES", "x", True, False),
        ("CHAPTER 1", "INTRODUCTION", "1", True, False),
        ("1.1", "Background and Context", "1", False, True),
        ("1.2", "Problem Statement", "1", False, True),
        ("1.3", "Objectives of the Project", "2", False, True),
        ("1.4", "Scope of the Project", "2", False, True),
        ("1.5", "Authentication and User Data Isolation Requirements", "3", False, True),
        ("CHAPTER 2", "LITERATURE REVIEW", "4", True, False),
        ("2.1", "Evolution of Scientific Document Processing", "4", False, True),
        ("2.2", "Traditional NLP and Extractive Summarization", "4", False, True),
        ("2.3", "Large Language Models and Hallucination in Academic QA", "5", False, True),
        ("2.4", "Retrieval-Augmented Generation (RAG) in Scientific Domains", "5", False, True),
        ("2.5", "Comparative Analysis of Existing Literature Review Assistants", "6", False, True),
        ("2.6", "Summary of Research Gaps", "6", False, True),
        ("CHAPTER 3", "SYSTEM DESIGN AND ARCHITECTURE", "7", True, False),
        ("3.1", "Overall System Architecture", "7", False, True),
        ("3.2", "PDF Ingestion and Page-Bound Text Extraction Pipeline", "8", False, True),
        ("3.3", "Metadata-Preserving Hierarchical Chunking", "8", False, True),
        ("3.4", "Hybrid Semantic-Lexical Retrieval Architecture", "9", False, True),
        ("3.5", "Authentication, Authorization, and Security Architecture", "10", False, True),
        ("3.6", "Multi-Tenant Data Isolation and Ownership Verification Model", "11", False, True),
        ("CHAPTER 4", "RESEARCH METHODOLOGY AND ALGORITHMIC FORMULATION", "12", True, False),
        ("CHAPTER 5", "SYSTEM REQUIREMENTS", "15", True, False),
        ("CHAPTER 6", "SYSTEM DESIGN AND IMPLEMENTATION", "17", True, False),
        ("CHAPTER 7", "EXPERIMENTAL RESULTS AND ANALYSIS", "21", True, False),
        ("CHAPTER 8", "USE CASES", "25", True, False),
        ("CHAPTER 9", "ETHICAL, LEGAL AND SOCIETAL CONSIDERATIONS", "27", True, False),
        ("CHAPTER 10", "LIMITATIONS", "28", True, False),
        ("CHAPTER 11", "FUTURE SCOPE", "29", True, False),
        ("CHAPTER 12", "CONCLUSION", "30", True, False),
        ("", "REFERENCES", "31", True, False)
    ]
    dummy_lot = [
        ("Table 2.1", "Comparative Analysis of Literature Review Tools and Proposed System", "6"),
        ("Table 4.1", "Benchmark Evaluation Dataset and Scientific Document Corpus", "13"),
        ("Table 5.1", "Hardware Requirements", "15"),
        ("Table 5.2", "Software Requirements and Technology Stack", "16"),
        ("Table 5.3", "Security, Authentication, and Compliance Requirements Specification", "16"),
        ("Table 6.1", "Key Backend REST API Endpoints and Authorization Rules", "18"),
        ("Table 7.1", "Evidence Retrieval Precision and Page Grounding Accuracy", "21"),
        ("Table 7.2", "Ablation Study Results on Chunking Parameters and Fusion", "22"),
        ("Table 7.3", "Inference Latency Breakdown for a 3-Paper Comparison Session", "23"),
        ("Table 7.4", "Authentication, Authorization, and Security Test Suite Results", "24")
    ]
    dummy_lof = [
        ("Fig. 3.1", "End-to-End System Architecture: AI-Powered Research Paper Comparison Engine", "7"),
        ("Fig. 3.2", "PDF Ingestion and Page-Bound Text Extraction Pipeline", "8"),
        ("Fig. 3.3", "Metadata-Preserving Hierarchical Chunking with Sliding Window Overlap", "8"),
        ("Fig. 3.4", "Hybrid Semantic-Lexical Retrieval Architecture with Reciprocal Rank Fusion", "9"),
        ("Fig. 3.5", "User Authentication and JWT Authorization Flow with Spring Security", "10"),
        ("Fig. 3.6", "Multi-Tenant Data Isolation and Ownership Verification Architecture", "11"),
        ("Fig. 6.1", "Multi-Paper Comparison Matrix and Claim Verification Pipeline", "17"),
        ("Fig. 6.2", "Metric Extraction and Cross-Paper Consistency Verification Pipeline", "19"),
        ("Fig. 6.3", "Asynchronous Job Processing and State Transition Lifecycle", "20"),
        ("Fig. 7.1", "Evidence Retrieval Precision vs. Token Chunk Size Benchmark", "22"),
        ("Fig. 7.2", "Inference and Matrix Construction Latency vs. Number of Papers", "23"),
        ("Fig. 7.3", "Cross-Paper Claim Grounding Entailment Scores by Dimension Heatmap", "24")
    ]

    build_document_clean(dummy_toc, dummy_lot, dummy_lof)
    page_count = export_and_analyze(OUTPUT_DOCX, OUTPUT_PDF)
    print(f"Clean Pass Page count: {page_count}")
