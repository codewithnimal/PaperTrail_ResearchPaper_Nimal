"""
Full document inspection - get content around paragraphs 130-260 (TOC area and beginning of chapters)
and check all content paragraph by paragraph for the full mapping.
"""
from docx import Document
from docx.oxml.ns import qn

doc_path = r"c:\Users\Nimal V\OneDrive\Desktop\nimalprojects\paperTrial\frontend\documentation\PaperTrail_Project_Documentation.docx"
doc = Document(doc_path)

print("=== PARAGRAPHS 130 to 260 ===")
for i, para in enumerate(doc.paragraphs[130:260]):
    idx = i + 130
    text = para.text.strip()
    style = para.style.name
    pb = ""
    if 'pageBreak' in para._element.xml or 'w:type="page"' in para._element.xml:
        pb = " [PB]"
    sectPr = " [SECT]" if 'sectPr' in para._element.xml else ""
    if text or pb or sectPr:
        print(f"[{idx:4d}] [{style:25s}]{pb}{sectPr} |{text[:100]}|")

print("\n=== LAST 50 PARAGRAPHS (end of document) ===")
total = len(doc.paragraphs)
for i, para in enumerate(doc.paragraphs[total-50:]):
    idx = i + total - 50
    text = para.text.strip()
    style = para.style.name
    pb = ""
    if 'pageBreak' in para._element.xml or 'w:type="page"' in para._element.xml:
        pb = " [PB]"
    sectPr = " [SECT]" if 'sectPr' in para._element.xml else ""
    print(f"[{idx:4d}] [{style:25s}]{pb}{sectPr} |{text[:100]}|")
