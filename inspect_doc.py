"""
Script to inspect the existing Word document structure.
"""
from docx import Document
from docx.oxml.ns import qn
import re

doc_path = r"c:\Users\Nimal V\OneDrive\Desktop\nimalprojects\paperTrial\frontend\documentation\PaperTrail_Project_Documentation.docx"
doc = Document(doc_path)

print(f"Number of sections: {len(doc.sections)}")
print(f"Number of paragraphs: {len(doc.paragraphs)}")
print(f"Number of tables: {len(doc.tables)}")

print("\n=== STYLES USED ===")
styles_used = set()
for para in doc.paragraphs:
    styles_used.add(para.style.name)
for s in sorted(styles_used):
    print(f"  {s}")

print("\n=== FIRST 150 PARAGRAPHS ===")
for i, para in enumerate(doc.paragraphs[:150]):
    text = para.text.strip()
    style = para.style.name
    if text:
        print(f"[{i:4d}] [{style:30s}] {text[:100]}")

print("\n=== SECTION BREAKS ===")
for i, para in enumerate(doc.paragraphs):
    xml = para._element.xml
    if 'sectPr' in xml:
        print(f"[{i:4d}] Section break found: {para.text[:80]}")

print("\n=== HEADERS/FOOTERS PER SECTION ===")
for sec_idx, section in enumerate(doc.sections):
    print(f"\n--- Section {sec_idx} ---")
    print(f"  Start type: {section.start_type}")
    print(f"  Different first page: {section.different_first_page_header_footer}")
    try:
        hdr = section.header
        if hdr and hdr.paragraphs:
            for p in hdr.paragraphs:
                if p.text.strip():
                    print(f"  Header text: {p.text[:100]}")
    except:
        pass
    try:
        ftr = section.footer
        if ftr and ftr.paragraphs:
            for p in ftr.paragraphs:
                if p.text.strip():
                    print(f"  Footer text: {p.text[:100]}")
    except:
        pass

print("\n=== CHAPTER HEADINGS ===")
for i, para in enumerate(doc.paragraphs):
    style = para.style.name
    text = para.text.strip()
    if ('Heading' in style or 'heading' in style.lower()) and text:
        print(f"[{i:4d}] [{style}] {text[:120]}")
    elif text.upper() == text and len(text) > 5 and len(text) < 100 and 'CHAPTER' in text.upper():
        print(f"[{i:4d}] [CAPS] {text[:120]}")
